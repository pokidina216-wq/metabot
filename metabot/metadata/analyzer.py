"""
Анализатор метаданных файлов.
Поддержка: изображения, видео, аудио, PDF, DOCX, XLSX, PPTX, архивы.
"""
from __future__ import annotations

import io
import logging
import zipfile
import tarfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Optional

logger = logging.getLogger(__name__)


@dataclass
class MetadataResult:
    """Результат анализа метаданных."""
    file_name: str
    file_size: int
    file_type: str
    mime_type: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    gps: Optional[Dict[str, float]] = None
    error: Optional[str] = None

    def format_text(self) -> str:
        """Форматированный вывод для Telegram."""
        lines = [
            f"📄 <b>Файл:</b> {self.file_name}",
            f"📦 <b>Размер:</b> {self._human_size(self.file_size)}",
            f"🔖 <b>Тип:</b> {self.file_type}",
            f"📋 <b>MIME:</b> {self.mime_type}",
            "",
        ]

        if self.error:
            lines.append(f"⚠️ <b>Ошибка:</b> {self.error}")
            return "\n".join(lines)

        if self.metadata:
            lines.append("📊 <b>Метаданные:</b>")
            for key, value in self.metadata.items():
                if value:
                    lines.append(f"  • <b>{key}:</b> {value}")

        if self.gps:
            lat = self.gps.get("latitude")
            lon = self.gps.get("longitude")
            lines.append("")
            lines.append("📍 <b>GPS координаты:</b>")
            lines.append(f"  • Широта: <code>{lat}</code>")
            lines.append(f"  • Долгота: <code>{lon}</code>")
            if self.gps.get("altitude") is not None:
                lines.append(f"  • Высота: <code>{self.gps['altitude']} м</code>")
            if lat is not None and lon is not None:
                g, y, o = self._map_links(lat, lon)
                lines.append(
                    f"  • 🗺 <a href=\"{g}\">Google Maps</a> · "
                    f"<a href=\"{y}\">Яндекс</a> · "
                    f"<a href=\"{o}\">OpenStreetMap</a>"
                )

        if not self.metadata and not self.gps:
            lines.append("ℹ️ Метаданные не найдены или были удалены.")

        return "\n".join(lines)

    @staticmethod
    def _human_size(size: int) -> str:
        for unit in ("Б", "КБ", "МБ", "ГБ"):
            if size < 1024:
                return f"{size:.1f} {unit}"
            size /= 1024
        return f"{size:.1f} ТБ"

    @staticmethod
    def _map_links(lat: float, lon: float) -> tuple[str, str, str]:
        """Ссылки на координаты в Google Maps, Яндекс.Картах и OpenStreetMap."""
        google = f"https://www.google.com/maps?q={lat},{lon}"
        yandex = f"https://yandex.ru/maps/?pt={lon},{lat}&z=17&l=map"
        osm = f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}#map=17/{lat}/{lon}"
        return google, yandex, osm

    def to_dict(self) -> Dict[str, Any]:
        """Структурированное представление для JSON-выгрузки."""
        data: Dict[str, Any] = {
            "file_name": self.file_name,
            "file_size_bytes": self.file_size,
            "file_size_human": self._human_size(self.file_size),
            "file_type": self.file_type,
            "mime_type": self.mime_type,
            "metadata": self.metadata,
        }
        if self.error:
            data["error"] = self.error
        if self.gps:
            lat = self.gps.get("latitude")
            lon = self.gps.get("longitude")
            gps = dict(self.gps)
            if lat is not None and lon is not None:
                g, y, o = self._map_links(lat, lon)
                gps["maps"] = {"google": g, "yandex": y, "openstreetmap": o}
            data["gps"] = gps
        return data

    def to_full_report(self) -> str:
        """Полная текстовая сводка (без обрезки) для выгрузки файлом."""
        from datetime import datetime, timezone

        out = [
            "==================================================",
            "          VEXIS — ПОЛНАЯ СВОДКА МЕТАДАННЫХ",
            "==================================================",
            f"Файл:   {self.file_name}",
            f"Размер: {self._human_size(self.file_size)} ({self.file_size} байт)",
            f"Тип:    {self.file_type}",
            f"MIME:   {self.mime_type}",
            f"Сформировано: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
            "--------------------------------------------------",
        ]
        if self.error:
            out.append(f"ОШИБКА: {self.error}")

        if self.metadata:
            out.append("")
            out.append("МЕТАДАННЫЕ:")
            for key, value in self.metadata.items():
                if value:
                    out.append(f"  {key}: {value}")

        if self.gps:
            lat = self.gps.get("latitude")
            lon = self.gps.get("longitude")
            out.append("")
            out.append("GPS-КООРДИНАТЫ:")
            out.append(f"  Широта:  {lat}")
            out.append(f"  Долгота: {lon}")
            if self.gps.get("altitude") is not None:
                out.append(f"  Высота:  {self.gps['altitude']} м")
            if lat is not None and lon is not None:
                g, y, o = self._map_links(lat, lon)
                out.append("  Карты:")
                out.append(f"    Google Maps:   {g}")
                out.append(f"    Яндекс.Карты:  {y}")
                out.append(f"    OpenStreetMap: {o}")

        if not self.metadata and not self.gps and not self.error:
            out.append("")
            out.append("Метаданные не найдены или были удалены.")

        out.append("")
        out.append("==================================================")
        out.append("Vexis — анализ метаданных. Данные извлечены из файла как есть.")
        return "\n".join(out)


class MetadataAnalyzer:
    """Главный класс анализа метаданных."""

    async def analyze(self, file_data: bytes, file_name: str, mime_type: str) -> MetadataResult:
        """Анализировать файл и вернуть метаданные."""
        file_type = self._detect_type(file_name, mime_type)
        result = MetadataResult(
            file_name=file_name,
            file_size=len(file_data),
            file_type=file_type,
            mime_type=mime_type,
        )

        try:
            if file_type == "image":
                await self._analyze_image(file_data, result)
            elif file_type == "video":
                await self._analyze_video(file_data, result)
            elif file_type == "audio":
                await self._analyze_audio(file_data, result)
            elif file_type == "pdf":
                await self._analyze_pdf(file_data, result)
            elif file_type == "docx":
                await self._analyze_docx(file_data, result)
            elif file_type == "xlsx":
                await self._analyze_xlsx(file_data, result)
            elif file_type == "pptx":
                await self._analyze_pptx(file_data, result)
            elif file_type == "archive":
                await self._analyze_archive(file_data, file_name, result)
            else:
                result.metadata["info"] = "Формат не поддерживает извлечение метаданных"
        except Exception as e:
            logger.exception("Metadata analysis error: %s", e)
            result.error = f"Ошибка анализа: {str(e)}"

        return result

    # ── Изображения ────────────────────────────────────────
    async def _analyze_image(self, data: bytes, result: MetadataResult) -> None:
        from PIL import Image
        from PIL.ExifTags import TAGS, GPSTAGS

        img = Image.open(io.BytesIO(data))
        result.metadata["Разрешение"] = f"{img.width}×{img.height}"
        result.metadata["Формат"] = img.format or "N/A"
        result.metadata["Режим"] = img.mode

        exif_data = img.getexif()
        if not exif_data:
            return

        for tag_id, value in exif_data.items():
            tag_name = TAGS.get(tag_id, str(tag_id))
            if tag_name == "GPSInfo":
                continue  # Обрабатываем отдельно
            if isinstance(value, bytes):
                continue
            result.metadata[tag_name] = str(value)[:200]

        # GPS
        gps_info = exif_data.get_ifd(0x8825)
        if gps_info:
            gps = self._parse_gps(gps_info)
            if gps:
                result.gps = gps

    @staticmethod
    def _parse_gps(gps_info: dict) -> Optional[Dict[str, float]]:
        """Парсинг GPS из EXIF."""
        try:
            def to_degrees(values):
                d, m, s = values
                return float(d) + float(m) / 60 + float(s) / 3600

            lat = to_degrees(gps_info.get(2, (0, 0, 0)))
            lon = to_degrees(gps_info.get(4, (0, 0, 0)))
            if gps_info.get(1) == "S":
                lat = -lat
            if gps_info.get(3) == "W":
                lon = -lon
            if lat == 0 and lon == 0:
                return None
            gps = {"latitude": round(lat, 6), "longitude": round(lon, 6)}
            # Высота (GPSAltitude=6, GPSAltitudeRef=5: 0=над, 1=под уровнем моря)
            alt = gps_info.get(6)
            if alt is not None:
                try:
                    alt_val = float(alt)
                    if gps_info.get(5) == 1:
                        alt_val = -alt_val
                    gps["altitude"] = round(alt_val, 1)
                except (TypeError, ValueError):
                    pass
            return gps
        except Exception:
            return None

    # ── Видео ──────────────────────────────────────────────
    async def _analyze_video(self, data: bytes, result: MetadataResult) -> None:
        import tempfile
        import subprocess
        import json

        with tempfile.NamedTemporaryFile(suffix=".mp4", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            try:
                proc = subprocess.run(
                    [
                        "ffprobe", "-v", "quiet",
                        "-print_format", "json",
                        "-show_format", "-show_streams",
                        tmp.name,
                    ],
                    capture_output=True, text=True, timeout=30,
                )
                if proc.returncode == 0:
                    info = json.loads(proc.stdout)
                    fmt = info.get("format", {})
                    result.metadata["Длительность"] = f"{float(fmt.get('duration', 0)):.1f} сек"
                    result.metadata["Битрейт"] = fmt.get("bit_rate", "N/A")
                    result.metadata["Формат"] = fmt.get("format_long_name", "N/A")

                    tags = fmt.get("tags", {})
                    if tags.get("creation_time"):
                        result.metadata["Дата создания"] = tags["creation_time"]

                    for stream in info.get("streams", []):
                        if stream.get("codec_type") == "video":
                            result.metadata["Видеокодек"] = stream.get("codec_long_name", "N/A")
                            result.metadata["Разрешение"] = (
                                f"{stream.get('width', '?')}×{stream.get('height', '?')}"
                            )
                            result.metadata["FPS"] = stream.get("r_frame_rate", "N/A")
                        elif stream.get("codec_type") == "audio":
                            result.metadata["Аудиокодек"] = stream.get("codec_long_name", "N/A")
                            result.metadata["Частота"] = f"{stream.get('sample_rate', 'N/A')} Hz"
            except FileNotFoundError:
                result.metadata["info"] = "ffprobe не установлен — видеоанализ недоступен"
            except subprocess.TimeoutExpired:
                result.error = "Таймаут анализа видео"

    # ── Аудио ──────────────────────────────────────────────
    async def _analyze_audio(self, data: bytes, result: MetadataResult) -> None:
        import tempfile
        from mutagen import File as MutagenFile

        with tempfile.NamedTemporaryFile(suffix=".mp3", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            audio = MutagenFile(tmp.name)
            if audio is None:
                result.metadata["info"] = "Не удалось прочитать аудио-метаданные"
                return

            if audio.info:
                result.metadata["Длительность"] = f"{audio.info.length:.1f} сек"
                if hasattr(audio.info, "bitrate"):
                    result.metadata["Битрейт"] = f"{audio.info.bitrate // 1000} kbps"
                if hasattr(audio.info, "sample_rate"):
                    result.metadata["Частота"] = f"{audio.info.sample_rate} Hz"
                if hasattr(audio.info, "channels"):
                    result.metadata["Каналы"] = str(audio.info.channels)

            if audio.tags:
                for key in ("title", "artist", "album", "date", "genre"):
                    val = audio.tags.get(key) or audio.tags.get(key.upper())
                    if val:
                        result.metadata[key.capitalize()] = str(val)[:200]

    # ── PDF ────────────────────────────────────────────────
    async def _analyze_pdf(self, data: bytes, result: MetadataResult) -> None:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        result.metadata["Страниц"] = str(len(reader.pages))

        info = reader.metadata
        if info:
            mapping = {
                "/Author": "Автор",
                "/Creator": "Создатель",
                "/Producer": "Программа",
                "/Title": "Заголовок",
                "/Subject": "Тема",
                "/CreationDate": "Дата создания",
                "/ModDate": "Дата изменения",
            }
            for pdf_key, label in mapping.items():
                val = info.get(pdf_key)
                if val:
                    result.metadata[label] = str(val)[:200]

    # ── DOCX ───────────────────────────────────────────────
    async def _analyze_docx(self, data: bytes, result: MetadataResult) -> None:
        from docx import Document

        doc = Document(io.BytesIO(data))
        props = doc.core_properties
        result.metadata["Автор"] = props.author or "N/A"
        result.metadata["Последний автор"] = props.last_modified_by or "N/A"
        result.metadata["Дата создания"] = str(props.created) if props.created else "N/A"
        result.metadata["Дата изменения"] = str(props.modified) if props.modified else "N/A"
        result.metadata["Заголовок"] = props.title or "N/A"
        result.metadata["Тема"] = props.subject or "N/A"
        result.metadata["Версия"] = props.revision or "N/A"
        result.metadata["Категория"] = props.category or "N/A"
        result.metadata["Параграфов"] = str(len(doc.paragraphs))
        result.metadata["Таблиц"] = str(len(doc.tables))

    # ── XLSX ───────────────────────────────────────────────
    async def _analyze_xlsx(self, data: bytes, result: MetadataResult) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True)
        props = wb.properties
        result.metadata["Автор"] = props.creator or "N/A"
        result.metadata["Последний автор"] = props.lastModifiedBy or "N/A"
        result.metadata["Дата создания"] = str(props.created) if props.created else "N/A"
        result.metadata["Дата изменения"] = str(props.modified) if props.modified else "N/A"
        result.metadata["Заголовок"] = props.title or "N/A"
        result.metadata["Листов"] = str(len(wb.sheetnames))
        result.metadata["Названия листов"] = ", ".join(wb.sheetnames[:10])
        wb.close()

    # ── PPTX ───────────────────────────────────────────────
    async def _analyze_pptx(self, data: bytes, result: MetadataResult) -> None:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        props = prs.core_properties
        result.metadata["Автор"] = props.author or "N/A"
        result.metadata["Последний автор"] = props.last_modified_by or "N/A"
        result.metadata["Дата создания"] = str(props.created) if props.created else "N/A"
        result.metadata["Дата изменения"] = str(props.modified) if props.modified else "N/A"
        result.metadata["Заголовок"] = props.title or "N/A"
        result.metadata["Слайдов"] = str(len(prs.slides))

    # ── Архивы ─────────────────────────────────────────────
    async def _analyze_archive(
        self, data: bytes, file_name: str, result: MetadataResult
    ) -> None:
        name_lower = file_name.lower()

        if name_lower.endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                files = zf.namelist()
                result.metadata["Файлов в архиве"] = str(len(files))
                total = sum(i.file_size for i in zf.infolist())
                result.metadata["Общий размер"] = MetadataResult._human_size(total)
                result.metadata["Содержимое"] = ", ".join(files[:20])
                if len(files) > 20:
                    result.metadata["Содержимое"] += f" ... (+{len(files) - 20})"
        elif name_lower.endswith((".tar", ".tar.gz", ".tgz", ".tar.bz2")):
            with tarfile.open(fileobj=io.BytesIO(data)) as tf:
                members = tf.getmembers()
                result.metadata["Файлов в архиве"] = str(len(members))
                names = [m.name for m in members[:20]]
                result.metadata["Содержимое"] = ", ".join(names)
        else:
            result.metadata["info"] = "Формат архива не поддерживается"

    # ── Утилиты ────────────────────────────────────────────
    @staticmethod
    def _detect_type(file_name: str, mime_type: str) -> str:
        """Определить тип файла."""
        ext = Path(file_name).suffix.lower()
        mime_lower = mime_type.lower()

        if ext in (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp", ".heic"):
            return "image"
        if mime_lower.startswith("image/"):
            return "image"
        if ext in (".mp4", ".avi", ".mkv", ".mov", ".wmv", ".flv", ".webm"):
            return "video"
        if mime_lower.startswith("video/"):
            return "video"
        if ext in (".mp3", ".wav", ".flac", ".ogg", ".aac", ".m4a", ".wma"):
            return "audio"
        if mime_lower.startswith("audio/"):
            return "audio"
        if ext == ".pdf" or mime_lower == "application/pdf":
            return "pdf"
        if ext == ".docx":
            return "docx"
        if ext == ".xlsx":
            return "xlsx"
        if ext == ".pptx":
            return "pptx"
        if ext in (".zip", ".tar", ".tar.gz", ".tgz", ".tar.bz2", ".rar", ".7z"):
            return "archive"
        return "unknown"
